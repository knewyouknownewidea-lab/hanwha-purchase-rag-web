"""브라우저 E2E 테스트 (v0.3.0)
간이 로그인·사용자별 저장, 브랜드 헤더, 작은 채점 위젯, 선택 정보(결과창), 2분할·스크롤 1개·페이지 나누기, 평가 근거 팝오버, 떠 있는 챗봇(왼쪽 아래),
결과 형식·보기 설정·분석 이력·엑셀, 필터(부품 분류)·정렬·관심 목록, 새 하위 보기(결제 일정·원자재·인건비·경쟁사·업계 동향),
업로드 저장 유지, 해상도별 화면을 확인한다.
실행: (서버 실행 후) python3 tests/e2e.py [http://localhost:5173] [스크린샷 폴더]
"""
import sys, os, json, tempfile
from playwright.sync_api import sync_playwright

URL = sys.argv[1] if len(sys.argv) > 1 else "http://localhost:5173"
SHOTS = sys.argv[2] if len(sys.argv) > 2 else tempfile.mkdtemp()
os.makedirs(SHOTS, exist_ok=True)
results = []


def check(name, cond, detail=""):
    results.append((name, bool(cond), detail))
    print(("PASS " if cond else "FAIL ") + name + (f" — {detail}" if detail else ""))


def open_chat(page):
    if not page.locator("#chat-input").is_visible():
        page.click(".chat-launcher")
        page.wait_for_selector("#chat-input")


def ask(page, text):
    open_chat(page)
    page.fill("#chat-input", text)
    page.press("#chat-input", "Enter")
    page.wait_for_function("() => !document.querySelector('.progress')", timeout=8000)


def close_chat(page):
    if page.locator("#chat-window").count():
        page.locator(".chat-tools button[aria-label='닫기']").click()


def rows(page, sel=".data-table tbody tr"):
    return page.locator(sel).count()


def count_text(page):
    return page.locator(".filter-count b").inner_text()


def demo_login(page):
    page.goto(URL, wait_until="domcontentloaded")
    page.wait_for_selector("#login-demo")
    page.click("#login-demo")
    page.wait_for_selector(".eval-card")  # 저장된 업로드를 읽은 뒤 화면과 첫 사용 안내가 함께 뜬다
    if page.locator("#ob-skip").count():
        page.click("#ob-skip")


def small_fonts(page):
    return page.evaluate(
        """() => { const out=[]; for (const el of document.querySelectorAll('body *')) { if (!el.childNodes.length) continue; const hasText=[...el.childNodes].some(n=>n.nodeType===3&&n.textContent.trim()); if(!hasText) continue; const r=el.getBoundingClientRect(); if(!r.width) continue; const fs=parseFloat(getComputedStyle(el).fontSize); if (fs<12) out.push(el.tagName+'.'+el.className+':'+fs); } return out.slice(0,5); }"""
    )


def v_scrollers(page, root=".panel-a"):
    return page.evaluate(
        """(root) => [...document.querySelectorAll(root+' *')].filter(e=>{const s=getComputedStyle(e);return /(auto|scroll)/.test(s.overflowY)&&e.scrollHeight>e.clientHeight+1}).map(e=>e.className)""",
        root,
    )


def sort_all(page, table_sel, cycle_none=True):
    """정렬할 수 있는 모든 열을 눌러 첫 방향 → 반대 방향 → (해제 또는 기본 순서)를 확인한다."""
    ths = page.locator(f"{table_sel} thead th:has(button.th-sort)")
    n = ths.count()
    bad = []
    for i in range(n):
        th = ths.nth(i)
        btn = th.locator("button.th-sort")
        name = btn.inner_text().strip()
        btn.click()
        a = th.get_attribute("aria-sort")
        up_on = th.locator(".sort-ic.is-asc").count() if a == "ascending" else th.locator(".sort-ic.is-desc").count()
        btn.click()
        b = th.get_attribute("aria-sort")
        btn.click()
        c = th.get_attribute("aria-sort")
        ok = {a, b} == {"ascending", "descending"} and up_on == 1 and th.locator(".sort-up").count() == 1 and th.locator(".sort-down").count() == 1
        if cycle_none:
            ok = ok and c == "none"
        if not ok:
            bad.append(f"{name}:{a}/{b}/{c}")
    return n, bad


def no_hscroll(page):
    return page.evaluate("() => document.documentElement.scrollWidth <= window.innerWidth")


with sync_playwright() as p:
    browser = p.chromium.launch()
    ctx = browser.new_context(viewport={"width": 1920, "height": 1080}, accept_downloads=True)
    page = ctx.new_page()
    errors = []
    page.on("pageerror", lambda e: errors.append(str(e)))
    # 샌드박스는 외부 글꼴(Google Fonts)을 받을 수 없어 그 연결 실패는 제외한다
    page.on("console", lambda m: errors.append(m.text) if m.type == "error" and "ERR_TUNNEL_CONNECTION_FAILED" not in m.text else None)

    # 0. v0.2.0 때 저장된 내용(사용자 구분 없음)을 넣어 두고 첫 로그인 사용자에게 옮겨지는지 본다
    page.goto(URL, wait_until="domcontentloaded")
    legacy = {"messages": [{"id": "old1", "role": "user", "text": "예전 버전에서 한 질문", "at": "2026-10-10T09:00:00.000Z"}], "results": {}, "view": {"resultId": None, "template": None}, "lastResultId": None, "pending": None}
    page.evaluate("(v) => localStorage.setItem('procure-ai:session.v2', JSON.stringify(v))", legacy)
    page.reload(wait_until="domcontentloaded")

    # 1. 간이 로그인
    page.wait_for_selector("#login-name")
    check("처음 열면 로그인 화면", page.locator(".login-card").is_visible() and page.locator(".eval-card").count() == 0)
    check("로그인 화면 안내: 시연용·보안 인증 아님", "보안 인증이 아닙니다" in page.locator(".login-note").inner_text())
    page.click("#login-submit")
    check("이름 없이 들어가기 → 안내 문구", page.locator("#login-err").count() == 1)
    check("안내 문구는 이름 칸 바로 아래 + 이름 칸에 초점", page.locator("label:has(#login-name) #login-err").count() == 1 and page.evaluate("() => document.activeElement && document.activeElement.id === 'login-name'"))
    page.fill("#login-name", "홍길동")
    page.fill("#login-dept", "구매팀")
    page.click("#login-submit")
    page.wait_for_selector(".onboarding")
    check("처음 로그인하면 사용 안내 3단계", page.locator(".ob-dots span").count() == 3)
    page.click("#ob-next")
    page.click("#ob-next")
    page.click("#ob-done")
    page.wait_for_selector(".eval-card")
    check("사용 안내 닫힘", page.locator(".onboarding").count() == 0)
    check("헤더에 사용자 이름", "홍길동" in page.locator(".user-btn").inner_text())
    check("v0.2.0 저장 내용 이전 후 예전 키 삭제", page.evaluate("() => localStorage.getItem('procure-ai:session.v2') === null"))

    # 2. 헤더 브랜드
    ko = page.locator(".brand-ko").bounding_box()
    en = page.locator(".brand-en").bounding_box()
    check("로고 표시", page.locator(".brand-logo").count() == 1 and page.locator(".brand-logo").evaluate("e => e.naturalWidth > 0"))
    check("'HANWHA ENGINE' 너비 = '한화엔진' 너비 (양끝 2px 이내)", abs(ko["x"] - en["x"]) <= 2 and abs((ko["x"] + ko["width"]) - (en["x"] + en["width"])) <= 2, f"{round(ko['width'])} / {round(en['width'])}")
    check("한화엔진이 플랫폼 이름보다 큰 글자", page.locator(".brand-ko").evaluate("e => parseFloat(getComputedStyle(e).fontSize)") > page.locator(".brand-product").evaluate("e => parseFloat(getComputedStyle(e).fontSize)"))
    check("헤더 흐름 문구 삭제", "AI 분석 요청 → 결과 형식 선택" not in page.locator(".topbar").inner_text())
    check("1920: 헤더 한 줄", page.locator(".topbar").bounding_box()["height"] < 72, str(round(page.locator(".topbar").bounding_box()["height"])))

    # 3. 2분할·글자 크기·스크롤
    a = page.locator(".panel-a").bounding_box()
    c = page.locator(".panel-c").bounding_box()
    ratio = round(a["width"] / (a["width"] + c["width"]) * 100)
    check("1920: 데이터 관리 : 분석 결과 = 45:55", abs(ratio - 45) <= 2, f"{ratio}:{100 - ratio}")
    check("1920: 가로 스크롤 없음", no_hscroll(page))
    check("가운데 챗봇 패널 없음(2분할)", page.locator(".panel-b").count() == 0 and page.locator(".splitter").count() == 1)
    sf = small_fonts(page)
    check("모든 글자 12px 이상", not sf, ", ".join(sf))
    vs = v_scrollers(page)
    check("왼쪽 패널 세로 스크롤 영역 1개", len(vs) == 1, ", ".join(vs))
    check("표가 옆으로 넘치면 오른쪽 가장자리 그림자", page.locator(".panel-a .scroll-shell.can-r").count() >= 1)
    page.locator(".panel-a .scroll-shell.can-r .table-scroll").first.evaluate("e => { e.scrollLeft = e.scrollWidth; e.dispatchEvent(new Event('scroll')); }")
    page.wait_for_timeout(100)
    check("끝까지 밀면 오른쪽 그림자 대신 고정 열 그림자", page.locator(".panel-a .scroll-shell.can-l").count() >= 1 and page.locator(".panel-a .data-table").first.evaluate("t => !t.closest('.scroll-shell').classList.contains('can-r')"))
    page.locator(".panel-a .scroll-shell.can-l .table-scroll").first.evaluate("e => { e.scrollLeft = 0; e.dispatchEvent(new Event('scroll')); }")
    page.screenshot(path=f"{SHOTS}/01_1920_initial.png")

    # 4. 요약 지표 버튼 제거 · 채점 카드는 그대로, 처리 단계만 오른쪽 위 구석에 작게
    check("요약 지표 버튼 없음", page.locator(".kpi").count() == 0 and page.locator(".kpi-strip").count() == 0)
    card = page.locator(".eval-card")
    check("채점 카드 유지 (도넛·항목 막대 5개)", card.locator(".donut").count() == 1 and card.locator(".eval-bars li").count() == 5)
    check("카드 아래 처리 단계 영역 없음", page.locator(".eval-progress").count() == 0)
    sm = card.locator(".eval-steps-mini")
    cb = card.bounding_box(); sb = sm.bounding_box()
    check("처리 단계 5개는 카드 오른쪽 위 구석에 작게", sm.locator(".step-mini").count() == 5 and sb["y"] < cb["y"] + 90 and sb["x"] + sb["width"] > cb["x"] + cb["width"] - 40 and sb["height"] <= 32, f"{round(sb['height'])}px")
    check("처리 단계 이름·상태는 마우스를 올리면 표시", "자료 업로드 · 완료" in sm.locator(".step-mini").first.get_attribute("title"))
    page.locator(".chip-btn[data-filter='gradeCD']").click()
    check("C·D등급 빠른 필터는 그대로 (2곳)", count_text(page) == "2")
    page.locator(".chip-btn[data-filter='gradeCD']").click()

    # 5. 데이터 탭 4등분
    widths = [page.locator(".tab").nth(i).bounding_box()["width"] for i in range(4)]
    check("데이터 탭 4등분", max(widths) - min(widths) < 2, ", ".join(str(round(w)) for w in widths))

    # 6. 경계선
    sp = page.locator(".splitter").bounding_box()
    page.mouse.move(sp["x"] + sp["width"] / 2, sp["y"] + sp["height"] / 2)
    page.mouse.down()
    page.mouse.move(sp["x"] + 200, sp["y"] + sp["height"] / 2, steps=5)
    page.mouse.up()
    a2 = page.locator(".panel-a").bounding_box()["width"]
    check("경계선 끌어서 너비 조절", a2 > a["width"] + 150, f"{round(a['width'])} → {round(a2)}px")
    page.reload(wait_until="domcontentloaded")
    page.wait_for_selector(".eval-card")
    check("새로고침해도 로그인 유지", page.locator(".login-card").count() == 0)
    a3 = page.locator(".panel-a").bounding_box()["width"]
    check("조절한 너비 새로고침 후 유지", abs(a3 - a2) < 4)
    page.locator(".splitter").dblclick()
    check("경계선 두 번 누르면 45:55 복귀", abs(page.locator(".panel-a").bounding_box()["width"] - a["width"]) < 4)

    # 7. 평가 근거 팝오버
    btn = page.locator(".eval-head .eval-basis-btn")
    check("카드 오른쪽 위 '근거' 버튼", btn.count() == 1 and "근거" in btn.inner_text())
    btn.click()
    page.wait_for_selector("#basis-popover")
    pop = page.locator("#basis-popover")
    check("팝오버: 항목별 계산 근거 5개·데이터 출처", pop.locator(".basis-list li").count() == 5 and "데이터 출처" in pop.inner_text())
    check("팝오버: 항목별 점수 막대와 등급·종합 점수", pop.locator(".basis-bar").count() == 5 and "A등급" in pop.locator("h3").inner_text())
    check("팝오버: 챗봇 리포트 버튼", pop.locator(".pop-foot button").count() == 1)
    check("팝오버 열면 안으로 초점", page.evaluate("() => document.activeElement && document.activeElement.id === 'basis-popover'"))
    page.keyboard.press("Escape")
    check("Esc로 닫히고 버튼으로 초점", page.locator("#basis-popover").count() == 0 and page.evaluate("() => document.activeElement && document.activeElement.classList.contains('eval-basis-btn')"))
    btn.click()
    page.mouse.click(1400, 700)
    check("바깥을 누르면 닫힘", page.locator("#basis-popover").count() == 0)

    # 8. 템플릿 바: 결과 전 비활성
    check("결과 전 템플릿 4종 표시·비활성", page.locator(".tpl-btn").count() == 4 and page.locator(".tpl-btn:disabled").count() == 4)
    labels = page.locator(".tpl-btn-label").all_inner_texts()
    check("템플릿 이름", labels == ["협력사 비교표", "데이터 시각화", "종합 평가 카드", "분석 보고서"], ", ".join(labels))

    # 9. 챗봇: 왼쪽 아래 아이콘
    check("챗봇은 닫힌 아이콘으로 시작", page.locator(".chat-launcher").is_visible() and page.locator("#chat-window").count() == 0)
    lb = page.locator(".chat-launcher").bounding_box()
    check("챗봇 아이콘은 왼쪽 아래", lb["x"] < 60 and lb["y"] + lb["height"] > 1080 - 60, f"x={round(lb['x'])}, y={round(lb['y'])}")
    check("아이콘에 'AI에게 질문하기' 글자", "AI에게 질문하기" in page.locator(".chat-launcher").inner_text())
    check("첫 방문 안내 말풍선", page.locator(".chat-nudge").count() == 1)
    page.click(".chat-launcher")
    page.wait_for_selector("#chat-input")
    page.wait_for_timeout(100)
    check("아이콘을 누르면 창 열림 + 입력칸 초점", page.evaluate("() => document.activeElement && document.activeElement.id === 'chat-input'"))
    check("창을 열면 안내 말풍선 사라짐", page.locator(".chat-nudge").count() == 0)
    wb = page.locator("#chat-window").bounding_box()
    check("창 기본 위치는 왼쪽 아래(아이콘 위)", wb["x"] < 60 and wb["y"] + wb["height"] < lb["y"] + 10, f"x={round(wb['x'])}")
    page.wait_for_timeout(300)  # 여는 효과(0.2초)가 끝난 뒤 한 번에 잰다
    geo = page.evaluate("() => { const i=document.querySelector('.chat-title-icon').getBoundingClientRect(); const s=document.querySelector('.chat-title-icon svg').getBoundingClientRect(); return {dx: Math.abs((i.x+i.width/2)-(s.x+s.width/2)), dy: Math.abs((i.y+i.height/2)-(s.y+s.height/2)), w: s.width} }")
    check("창 머리 아이콘이 상자 가운데", geo["dx"] < 1.5 and geo["dy"] < 1.5 and geo["w"] >= 16, json.dumps({k: round(v, 1) for k, v in geo.items()}))
    check("창 머리 아이콘 흰색", page.locator(".chat-title-icon").evaluate("e => getComputedStyle(e).color") == "rgb(255, 255, 255)")
    page.keyboard.press("Escape")
    check("Esc로 닫힘", page.locator("#chat-window").count() == 0)
    page.click(".chat-launcher")
    page.wait_for_selector("#chat-window")

    # 10. 이동·크기·투명도(⋯ 메뉴) + 기억
    w0 = page.locator("#chat-window").bounding_box()
    hb = page.locator(".chat-title").bounding_box()
    page.mouse.move(hb["x"] + 20, hb["y"] + 10)
    page.mouse.down()
    page.mouse.move(hb["x"] + 420, hb["y"] - 60, steps=6)
    page.mouse.up()
    w1 = page.locator("#chat-window").bounding_box()
    check("머리를 끌어 창 이동", w1["x"] > w0["x"] + 300, f"x {round(w0['x'])} → {round(w1['x'])}")
    se = page.locator(".rz-se").bounding_box()
    page.mouse.move(se["x"] + 8, se["y"] + 8)
    page.mouse.down()
    page.mouse.move(se["x"] + 120, se["y"] + 8, steps=4)
    page.mouse.up()
    w2 = page.locator("#chat-window").bounding_box()
    check("모서리를 끌어 크기 조절", w2["width"] > w1["width"] + 80, f"w {round(w1['width'])} → {round(w2['width'])}")
    se = page.locator(".rz-se").bounding_box()
    page.mouse.move(se["x"] + 8, se["y"] + 8)
    page.mouse.down()
    page.mouse.move(se["x"] - 900, se["y"] - 900, steps=4)
    page.mouse.up()
    w3 = page.locator("#chat-window").bounding_box()
    check("최소 크기 320×400 제한", w3["width"] >= 319 and w3["height"] >= 399, f"{round(w3['width'])}×{round(w3['height'])}")
    check("머리에는 제목·아이콘만, 설정은 ⋯ 메뉴", page.locator("#chat-opacity").count() == 0 and page.locator(".chat-menu-btn").count() == 1)
    page.click(".chat-menu-btn")
    page.locator("#chat-opacity").fill("60")
    op = page.locator("#chat-window").evaluate("e => getComputedStyle(e).opacity")
    check("투명도 슬라이더(창 전체)", abs(float(op) - 0.6) < 0.01, op)
    page.keyboard.press("Escape")
    check("Esc는 메뉴부터 닫음", page.locator(".chat-menu").count() == 0 and page.locator("#chat-window").count() == 1)
    page.reload(wait_until="domcontentloaded")
    page.wait_for_selector(".eval-card")
    page.click(".chat-launcher")
    page.wait_for_selector("#chat-window")
    w4 = page.locator("#chat-window").bounding_box()
    op2 = page.locator("#chat-window").evaluate("e => getComputedStyle(e).opacity")
    check("창 위치·크기·투명도 새로고침 후 유지", abs(w4["x"] - w3["x"]) < 3 and abs(float(op2) - 0.6) < 0.01)
    page.click(".chat-menu-btn")
    page.locator(".chat-menu button[aria-label='위치·크기·투명도 초기화']").click()
    check("위치·크기·투명도 초기화", page.locator("#chat-window").evaluate("e => getComputedStyle(e).opacity") == "1" and page.locator("#chat-window").bounding_box()["x"] < 60)
    page.locator(".chat-tools button[aria-label='최소화']").click()
    check("최소화", page.locator("#chat-input").count() == 0)
    page.locator(".chat-tools button[aria-label='펼치기']").click()
    check("이전 사용자 대화 이전됨(예전 버전 질문)", "예전 버전에서 한 질문" in page.locator(".chat-list").inner_text())
    check("입력창 위 추천 질문 칩", page.locator(".suggest-row .suggest-chip").count() >= 4)
    check("입력 안내(Enter·Shift+Enter)", "Shift+Enter" in page.locator(".composer-hint").inner_text())

    # 11. 대표 시나리오
    q = "해당 품목의 과거 거래 업체를 등급, 납기 실적, 최근 단가와 함께 보여주어 품목별 후보 업체 표시해줘."
    page.click(".chat-menu-btn")
    page.locator(".chat-menu button[aria-label='대화 지우기']").click()
    check("대화 지우기 전 확인 창", page.locator("[role='alertdialog']").count() == 1 and "분석 이력" in page.locator("[role='alertdialog']").inner_text())
    page.keyboard.press("Escape")
    st = (page.locator("[role='alertdialog']").count(), page.locator("#chat-window").count(), "예전 버전에서 한 질문" in page.locator(".chat-list").inner_text())
    check("확인 창 Esc → 취소 (대화창·대화 그대로)", st == (0, 1, True), str(st))
    page.click(".chat-menu-btn")
    page.locator(".chat-menu button[aria-label='대화 지우기']").click()
    page.locator("[role='alertdialog'] [data-confirm]").click()
    check("대화 지우기 → 입력칸으로 초점", page.evaluate("() => new Promise(r => setTimeout(() => r(document.activeElement && document.activeElement.id === 'chat-input'), 50))"))
    page.locator(".ctx-chips .chip-x").first.click() if page.locator(".ctx-chips .chip-accent").count() else None
    ask(page, q)
    last = page.locator(".msg-ai").last
    check("품목 미선택 → 품목 되묻기", last.locator(".opt-chip").count() >= 5)
    check("되묻는 단계에서는 결과 미생성", page.locator(".empty-title").count() == 1)
    last.locator(".opt-chip", has_text="실린더 라이너").click()
    page.wait_for_function("() => !document.querySelector('.progress')")
    page.wait_for_selector(".result-table")
    check("분석 후 추천 형식으로 바로 표시", page.locator(".tpl-btn.is-active .tpl-btn-label").inner_text() == "협력사 비교표")
    check("추천 표시", page.locator(".tpl-btn.is-active .tpl-rec").count() == 1)
    check("대화에는 형식 선택 카드 대신 안내", page.locator(".msg-ai").last.locator(".tpl-card").count() == 0 and "분석 결과창 상단에서 선택" in page.locator(".msg-ai").last.inner_text())
    check("C 비교표 4개 업체 · 조건 상이 1", page.locator(".result-table tbody tr").count() == 4 and page.locator(".result-table tr.is-excluded").count() == 1)
    body = page.locator(".result").inner_text()
    for label in ["단가 경쟁력", "납기 실적", "추가 검토가 필요한 위험 요소", "비교에 사용된 데이터와 평가 근거", "자동 확정되지 않습니다"]:
        check(f"결과에 '{label}'", label in body)
    close_chat(page)
    page.screenshot(path=f"{SHOTS}/02_1920_table.png")

    # 12. 결과표 정렬 두 방향
    th = page.locator(".result-table thead th", has_text="납기 준수율").locator("button")
    th.click()
    s1 = page.locator(".result-table thead th", has_text="납기 준수율").get_attribute("aria-sort")
    th.click()
    s2 = page.locator(".result-table thead th", has_text="납기 준수율").get_attribute("aria-sort")
    check("비교표: 오름·내림 두 방향 정렬", {s1, s2} == {"ascending", "descending"}, f"{s1} → {s2}")
    th.click()
    check("비교표: 세 번째 누르면 기본 순서(검토 순서)로", page.locator(".result-table thead th.rank-col").get_attribute("aria-sort") == "ascending")
    n_s, bad_s = sort_all(page, ".result-table", cycle_none=False)
    check(f"비교표: 정렬할 수 있는 모든 열({n_s}개) 두 방향 + 위·아래 화살표", n_s >= 5 and not bad_s, ", ".join(bad_s))

    # 13. 템플릿 전환 (재분석 없음) + 보고서 가운데
    title = page.locator(".result-titles h3").inner_text()
    for tid, sel in [("chart", ".chart-card"), ("cards", ".score-card"), ("report", ".report-doc"), ("table", ".result-table")]:
        page.locator(f".tpl-btn[data-template='{tid}']").click()
        page.wait_for_selector(sel)
        check(f"상단 템플릿 전환: {tid}", page.locator(sel).count() >= 1)
        if tid == "report":
            rb = page.locator(".report-doc").bounding_box()
            body_box = page.locator(".result-body").bounding_box()
            check("분석 보고서 가운데 정렬", abs((rb["x"] + rb["width"] / 2) - (body_box["x"] + body_box["width"] / 2)) < 12, f"{round(rb['x'])}")
    check("템플릿만 바꿔도 같은 결과", page.locator(".result-titles h3").inner_text() == title)

    # 14. 보기 설정·프리셋
    page.click(".vs-toggle")
    page.wait_for_selector(".vs-panel")
    check("보기 설정 열림 · 기본 프리셋 4종", page.locator(".preset").count() == 4)
    page.locator(".preset[data-preset='delivery']").click()
    # 프리셋의 정렬은 useEffect에서 반영되므로 변경된 행이 렌더될 때까지 기다린다.
    page.wait_for_function("() => document.querySelector('.result-table tbody tr .sticky-col .name')?.textContent === '업체 D'")
    first = page.locator(".result-table tbody tr").first.locator(".sticky-col .name").inner_text()
    heads = page.locator(".result-table thead th").all_inner_texts()
    check("프리셋 '납기·공급 리스크': 납기 낮은 순 정렬", first == "업체 D", first)
    check("프리셋 '납기·공급 리스크': 지연 건수 열 표시", any("지연 건수" in h for h in heads))
    check("현재 프리셋 이름 표시", "납기·공급 리스크" in page.locator(".vs-current").inner_text())
    page.locator(".vs-panel .check", has_text="최근 단가").locator("input").uncheck()
    check("열 선택 해제", not any("최근 단가" in h for h in page.locator(".result-table thead th").all_inner_texts()))
    page.locator(".seg-ctl button", has_text="촘촘").click()
    check("행 간격 촘촘", "is-compact" in page.locator(".result-table").get_attribute("class"))
    page.fill("#preset-name", "내 협상용")
    page.locator(".preset-save button[type='submit']").click()
    check("내 프리셋 저장", page.locator(".user-preset", has_text="내 협상용").count() == 1)
    page.locator(".preset-save button", has_text="기본값으로").click()
    check("기본값으로 되돌리기", page.locator(".vs-current").inner_text() == "기본" and "is-compact" not in page.locator(".result-table").get_attribute("class"))
    page.locator(".user-preset", has_text="내 협상용").locator("button").first.click()
    check("내 프리셋 적용", "내 협상용" in page.locator(".vs-current").inner_text() and "is-compact" in page.locator(".result-table").get_attribute("class"))
    page.mouse.click(400, 40)
    page.locator(".tpl-btn[data-template='chart']").click()
    page.click(".vs-toggle")
    page.locator(".preset-save button", has_text="기본값으로").click()
    n_before = page.locator(".chart-card").count()
    page.locator(".vs-panel .check-grid .check").first.locator("input").uncheck()
    check("차트 숨기기", page.locator(".chart-card").count() == n_before - 1, f"{n_before} → {page.locator('.chart-card').count()}")
    page.locator(".seg-ctl button", has_text="1개").click()
    check("한 줄에 차트 1개", "per-1" in page.locator(".chart-grid").get_attribute("class"))
    page.mouse.click(400, 40)
    page.locator(".tpl-btn[data-template='cards']").click()
    page.click(".vs-toggle")
    page.locator(".vs-panel .seg-ctl button", has_text="3개").click()
    check("한 줄에 카드 3개", "per-3" in page.locator(".score-cards").get_attribute("class"))
    page.select_option("#vs-cards-sort", "onTimeRate")
    page.locator(".vs-panel .seg-ctl[aria-label='카드 정렬 방향'] button", has_text="오름차순").click()
    asc_first = page.locator(".score-card h4").first.inner_text()
    page.locator(".vs-panel .seg-ctl[aria-label='카드 정렬 방향'] button", has_text="내림차순").click()
    desc_first = page.locator(".score-card h4").first.inner_text()
    check("카드 정렬 방향(오름·내림) 선택", asc_first == "업체 D" and desc_first == "업체 A", f"{asc_first} / {desc_first}")
    page.mouse.click(400, 40)
    page.locator(".tpl-btn[data-template='report']").click()
    page.click(".vs-toggle")
    page.locator(".vs-panel .seg-ctl button", has_text="요약형").click()
    page.locator(".vs-panel .check", has_text="평가 근거·참조 데이터").locator("input").uncheck()
    check("보고서 요약형·섹션 선택", "is-brief" in page.locator(".report-doc").get_attribute("class") and page.locator(".report-doc section[data-section='basis']").count() == 0)
    page.screenshot(path=f"{SHOTS}/03_1920_settings.png")
    page.locator(".preset-save button", has_text="기본값으로").click()
    page.mouse.click(400, 40)
    page.locator(".tpl-btn[data-template='table']").click()

    # 15. 후속 질문
    ask(page, "여기서 단가보다 납기를 우선하면 어떤 업체가 적합해?")
    page.wait_for_function("() => document.querySelector('.result-titles h3').textContent.includes('납기 우선')")
    check("후속 질문: 같은 품목·납기 우선으로 바로 표시", "실린더 라이너" in page.locator(".result-titles h3").inner_text())
    check("후속 질문: 보던 형식 유지", page.locator(".tpl-btn.is-active .tpl-btn-label").inner_text() == "협력사 비교표")
    check("대화에 '보던 형식' 안내", "보던 형식" in page.locator(".msg-ai").last.inner_text())
    prev_btn = page.locator(".msg-ai .msg-result button", has_text="이 결과 다시 보기").first
    check("이전 결과 다시 보기 버튼", prev_btn.count() == 1)
    prev_btn.click()
    check("이전 결과로 전환", "납기 우선" not in page.locator(".result-titles h3").inner_text())

    # 16. 새로고침해도 대화·결과 유지
    n_msgs = page.locator(".msg").count()
    shown_title = page.locator(".result-titles h3").inner_text()
    page.reload(wait_until="domcontentloaded")
    page.wait_for_selector(".result-titles h3")
    check("새로고침 후 분석 결과 유지", page.locator(".result-titles h3").inner_text() == shown_title)
    open_chat(page)
    check("새로고침 후 대화 내역 유지", page.locator(".msg").count() == n_msgs, f"{n_msgs}개")

    # 17. 다른 질문
    for qq, tpl in [("납기 준수율이 낮은 협력사를 찾아줘.", "데이터 시각화"), ("최근 단가가 상승한 품목을 보여줘.", "협력사 비교표"), ("A등급 협력사의 품질 실적을 정리해 줘.", "종합 평가 카드")]:
        ask(page, qq)
        check(f"질문 → 바로 표시: {qq}", page.locator(".tpl-btn.is-active .tpl-btn-label").inner_text() == tpl)
    ask(page, "오늘 점심 메뉴 추천해줘")
    check("범위 밖 질문은 추측하지 않음", "추측해서 답하지 않습니다" in page.locator(".msg-ai").last.inner_text())

    # 18. 키워드 검색 → 데이터 관리 선택 (다른 쪽에 있어도 그 쪽으로)
    page.fill("#keyword-search", "PO-2024-0003")
    page.wait_for_selector(".kw-hit")
    page.locator(".kw-hit").first.click()
    check("키워드 검색 → 데이터 관리에서 선택 (페이지 이동)", page.locator(".data-table tr.is-selected").count() == 1)
    close_chat(page)

    # 19. 분석 이력
    page.click(".history-toggle")
    page.wait_for_selector("#history-panel")
    n_hist = page.locator(".history-item").count()
    check("분석 이력 목록 (이번 세션 결과들)", n_hist >= 5, str(n_hist))
    page.fill("#history-q", "납기 우선")
    check("이력 검색", page.locator(".history-item").count() == 1)
    page.locator(".history-item").first.locator("button[title='이름 바꾸기']").click()
    page.fill("#history-rename-input", "라이너 납기 우선 검토")
    page.press("#history-rename-input", "Enter")
    check("이름 바꾸기", "라이너 납기 우선 검토" in page.locator(".history-item").first.inner_text())
    page.locator(".history-item").first.locator("button[title='고정 (자동 삭제에서 제외)']").click()
    check("고정", page.locator(".history-item").first.locator("svg[aria-label='고정됨']").count() == 1)
    page.locator(".history-item .history-open").first.click()
    check("이력에서 누르면 결과창에 표시 (바꾼 이름)", page.locator(".result-titles h3").inner_text() == "라이너 납기 우선 검토")
    page.click(".history-toggle")
    page.fill("#history-q", "점심")
    check("없는 검색어 안내", page.locator(".history-item").count() == 0)
    page.fill("#history-q", "")
    before = page.locator(".history-item").count()
    page.locator(".history-item").last.locator("button[title='삭제']").click()
    check("이력 삭제 전 한 번 더 묻기", page.locator(".history-confirm").count() == 1 and page.locator(".history-item").count() == before)
    page.locator(".history-confirm .history-del-yes").click()
    check("이력 삭제", page.locator(".history-item").count() == before - 1)
    page.keyboard.press("Escape")

    # 20. 엑셀(.xlsx) 내보내기
    import openpyxl
    page.click(".export-toggle")
    with page.expect_download() as dl:
        page.locator(".menu-item[data-export='xlsx']").click()
    d1 = dl.value
    x1 = os.path.join(SHOTS, d1.suggested_filename)
    d1.save_as(x1)
    wb = openpyxl.load_workbook(x1)
    check("결과 엑셀: 파일 이름 분석결과_…_YYYYMMDD.xlsx", d1.suggested_filename.startswith("분석결과_") and d1.suggested_filename.endswith(".xlsx"), d1.suggested_filename)
    check("결과 엑셀: 시트 '결과 표'·'조건·근거'", wb.sheetnames == ["결과 표", "조건·근거"], ",".join(wb.sheetnames))
    ws = wb["결과 표"]
    hdr = [c.value for c in ws[1]]
    check("결과 엑셀: 머리글 굵게·틀 고정·숫자 칸", ws["A1"].font.b and ws.freeze_panes == "A2" and any(isinstance(c.value, (int, float)) for c in ws[2]), str(hdr[:4]))
    check("결과 엑셀: SAMPLE 표기", any("SAMPLE" in str(r[1].value) for r in wb["조건·근거"].iter_rows()))
    page.locator(".tab", has_text="협력업체").click()
    page.locator(".active-chips .link-btn").click() if page.locator(".active-chips .link-btn").count() else None
    with page.expect_download() as dl:
        page.locator(".filter-row button[data-export='xlsx']").click()
    d2 = dl.value
    x2 = os.path.join(SHOTS, d2.suggested_filename)
    d2.save_as(x2)
    wb2 = openpyxl.load_workbook(x2)
    check("데이터 탭 엑셀: 지금 조건의 행 (협력업체 12행)", wb2.sheetnames[0] == "협력업체" and wb2["협력업체"].max_row == 13, f"{wb2.sheetnames} {wb2[wb2.sheetnames[0]].max_row}")

    # 21. 표 정렬 두 방향 + 페이지 나누기
    hdr_on = page.locator(".data-table thead th[data-col='ontime']")
    hdr_on.locator("button").click()
    s1 = hdr_on.get_attribute("aria-sort")
    top1 = page.locator(".data-table tbody tr").first.locator(".name").inner_text()
    hdr_on.locator("button").click()
    s2 = hdr_on.get_attribute("aria-sort")
    top2 = page.locator(".data-table tbody tr").first.locator(".name").inner_text()
    hdr_on.locator("button").click()
    s3 = hdr_on.get_attribute("aria-sort")
    check("숫자 열: 내림 → 오름 → 해제", (s1, s2, s3) == ("descending", "ascending", "none"), f"{s1}/{s2}/{s3} · {top1} → {top2}")
    check("숫자 열: 오름차순이면 가장 낮은 값이 먼저", top2 in ("업체 J", "업체 D"), top2)
    hdr_name = page.locator(".data-table thead th[data-col='name']")
    hdr_name.locator("button").click()
    check("글자 열: 첫 클릭 오름차순", hdr_name.get_attribute("aria-sort") == "ascending")
    hdr_name.locator("button").click()
    check("글자 열: 두 번째 내림차순", hdr_name.get_attribute("aria-sort") == "descending" and page.locator(".data-table tbody tr").first.locator(".name").inner_text() == "업체 L")
    check("정렬 아이콘: 위·아래 화살표 함께", page.locator(".data-table thead .sort-ic .sort-up").count() > 0 and page.locator(".data-table thead .sort-ic .sort-down").count() > 0)
    page.locator(".tab", has_text="구매조달").click()
    page.locator(".subtab[data-sub='list']").click()
    page.locator(".active-chips .link-btn").click() if page.locator(".active-chips .link-btn").count() else None
    check("선택한 행이 있는 쪽을 그대로 보여 줌 (검색으로 고른 PO-2024-0003)", page.locator(".data-table tr.is-selected", has_text="PO-2024-0003").count() == 1, page.locator(".pager-range").inner_text())
    page.locator(".row-detail .detail-clear").click()
    page.locator(".chip-btn[data-filter='thisMonth']").click()
    page.locator(".chip-btn[data-filter='thisMonth']").click()
    check("페이지 나누기: 기본 20행", rows(page) == 20 and "1–20 / 346행" in page.locator(".pager-range").inner_text(), page.locator(".pager-range").inner_text())
    page.locator(".pager button[aria-label='다음 쪽']").click()
    check("다음 쪽", "21–40" in page.locator(".pager-range").inner_text())
    page.select_option(".pager-size select", "50")
    check("한 쪽 50행", rows(page) == 50)
    vs2 = v_scrollers(page)
    check("구매조달 50행에서도 세로 스크롤 1개", len(vs2) == 1, ", ".join(vs2))
    hdr_due = page.locator(".data-table thead th[data-col='due']")
    hdr_due.locator("button").click()
    check("날짜 열: 첫 클릭 최근 순(내림차순)", hdr_due.get_attribute("aria-sort") == "descending")
    hdr_due.locator("button").click()
    hdr_due.locator("button").click()
    for tab_name in ["협력업체", "구매조달", "원가-계약", "자재-공급관리"]:
        page.locator(".tab", has_text=tab_name).click()
        if page.locator(".subtab[data-sub='list']").count():
            page.locator(".subtab[data-sub='list']").click()
        n_s, bad_s = sort_all(page, ".data-table:not(.mat-table):not(.link-table)")
        check(f"{tab_name}: 정렬할 수 있는 모든 열({n_s}개) 두 방향 → 해제", n_s >= 5 and not bad_s, ", ".join(bad_s))

    # 22. 필터 세분화 + 부품 분류
    page.locator(".tab", has_text="협력업체").click()
    page.locator(".chip-btn[data-filter='lowDelivery']").click()
    check("빠른 필터: 납기 90% 미만", count_text(page) == "4" and rows(page) == 4)
    check("적용 조건 칩 표시", page.locator(".fchip").count() == 1)
    page.locator(".filter-toggle").click()
    page.locator(".fd-multi .chip-btn", has_text="B등급").click()
    check("상세 필터(등급 다중 선택) 조합", int(count_text(page)) < 4 and page.locator(".fchip").count() == 2, count_text(page))
    page.locator("#f-suppliers-quality-min").fill("82")
    check("범위 필터(품질 점수)", page.locator(".fchip", has_text="품질 점수").count() == 1)
    page.locator(".fchip").first.locator("button").click()
    check("칩 × 로 조건 해제", page.locator(".fchip").count() == 2)
    page.locator(".active-chips .link-btn", has_text="전체 초기화").click()
    check("전체 초기화", count_text(page) == "12" and page.locator(".fchip").count() == 0)
    check("부품 분류: 앞 단계 전에는 다음 단계 비활성", page.locator("#f-suppliers-bom-1").is_disabled())
    page.select_option("#f-suppliers-bom-0", "SAMPLE 엔진 X형")
    page.select_option("#f-suppliers-bom-1", "실린더부 (라이너·커버·배기밸브)")
    opts = page.locator("#f-suppliers-bom-2 option").all_inner_texts()
    check("부품 분류: 위 단계를 고르면 아래 목록이 좁혀짐", "배기밸브 어셈블리" in opts and not any("피스톤" in o for o in opts), ", ".join(opts))
    page.select_option("#f-suppliers-bom-2", "배기밸브 어셈블리")
    names = sorted(page.locator(".data-table tbody .name").all_inner_texts())
    check("부품 분류: 배기밸브 스핀들 공급 협력사만 (업체 G·L)", names == ["업체 G", "업체 L"], ", ".join(names))
    check("부품 분류 조건 칩 (경로 ›)", page.locator(".fchip", has_text="›").count() == 1)
    page.screenshot(path=f"{SHOTS}/04_1920_filters.png")
    page.locator(".active-chips .link-btn").click()
    page.locator(".tab", has_text="구매조달").click()
    page.locator(".filter-toggle").click()
    page.select_option("#f-purchaseOrders-bom-0", "SAMPLE 엔진 Y형")
    page.select_option("#f-purchaseOrders-bom-1", "연료·유압 공급부")
    page.select_option("#f-purchaseOrders-bom-2", "고압 연료 배관")
    check("구매조달에도 부품 분류 필터 (고압 연료관 발주만)", set(page.locator(".data-table tbody .name-cell .name").all_inner_texts()) == {"고압 연료관"})
    page.locator(".active-chips .link-btn").click()
    page.locator(".chip-btn[data-filter='overdue']").click()
    n_over = int(count_text(page))
    check("구매조달: 납기 지연(미입고)", n_over == 3, str(n_over))
    page.locator(".chip-btn[data-filter='overdue']").click()
    page.locator("#f-purchaseOrders-due-from").fill("2026-10-01")
    page.locator("#f-purchaseOrders-due-to").fill("2026-10-31")
    n_month = int(count_text(page))
    page.locator(".active-chips .link-btn").click()
    page.locator(".chip-btn[data-filter='thisMonth']").click()
    check("기간 필터 = 이번 달 납기 빠른 필터", n_month == int(count_text(page)) and n_month > 0, str(n_month))
    page.locator(".active-chips .link-btn").click()
    page.locator(".tab", has_text="원가-계약").click()
    page.locator(".subtab[data-sub='list']").click()
    page.locator(".chip-btn[data-filter='up']").click()
    check("원가-계약: 5% 이상 상승", count_text(page) == "4")
    page.locator(".tab", has_text="자재-공급관리").click()
    page.locator(".subtab[data-sub='list']").click()
    page.locator(".active-chips .link-btn").click() if page.locator(".active-chips .link-btn").count() else None
    page.locator(".chip-btn[data-filter='below']").click()
    check("자재-공급: 안전재고 미달", count_text(page) == "3")
    page.locator(".tab", has_text="원가-계약").click()
    check("탭별 필터 따로 유지", count_text(page) == "4")
    page.locator(".active-chips .link-btn").click()
    page.locator(".tab", has_text="자재-공급관리").click()
    page.locator(".active-chips .link-btn").click()

    # 23. 관심 협력사·부품
    page.locator(".tab", has_text="협력업체").click()
    page.locator(".data-table tbody tr", has_text="업체 F").locator(".fav-star").click()
    check("별표로 관심 협력사 고정", page.locator(".data-table tbody tr", has_text="업체 F").locator(".fav-star").get_attribute("aria-pressed") == "true")
    page.locator(".chip-btn[data-filter='fav']").click()
    check("빠른 필터 '★ 관심만'", count_text(page) == "1")
    page.locator(".chip-btn[data-filter='fav']").click()
    page.check("#pin-fav")
    check("관심 항목 맨 위 고정", page.locator(".data-table tbody tr").first.locator(".name").inner_text() == "업체 F")
    page.locator(".tab", has_text="자재-공급관리").click()
    page.locator(".data-table tbody tr", has_text="터보차저 노즐링").locator(".fav-star").click()
    check("관심 부품 고정", page.locator(".data-table tbody tr", has_text="터보차저 노즐링").locator(".fav-star").get_attribute("aria-pressed") == "true")
    page.reload(wait_until="domcontentloaded")
    page.wait_for_selector(".eval-card")
    check("평가 대상 목록 맨 위에 관심 협력사", page.locator("#eval-supplier optgroup").first.get_attribute("label") == "★ 관심 협력사")
    open_chat(page)
    sugg = page.locator(".suggest-row .suggest-chip").all_inner_texts()
    check("추천 질문에 관심 협력사·부품", any("업체 F" in s for s in sugg) and any("터보차저 노즐링" in s for s in sugg), " | ".join(sugg[:2]))
    close_chat(page)
    page.click(".user-btn")
    check("사용자 메뉴에 관심 목록", "업체 F" in page.locator(".user-pop").inner_text() and "터보차저 노즐링" in page.locator(".user-pop").inner_text())
    page.keyboard.press("Escape")

    # 24. 협력사 상세: 공급 부품 분류 경로 · 등급 변동 이력 · 관련 직종 임금
    page.locator(".tab", has_text="협력업체").click()
    page.locator(".data-table tbody tr", has_text="업체 F").locator(".name").click()
    det = page.locator(".row-detail")
    check("행을 누르면 상세는 오른쪽 분석 결과창에", page.locator(".panel-c #selection-detail").count() == 1 and page.locator(".panel-a .row-detail").count() == 0)
    check("결과창 머리에 선택 정보 아이콘(펼침)", page.locator(".detail-toggle").get_attribute("aria-expanded") == "true" and "업체 F" in page.locator(".detail-toggle").inner_text())
    check("상세: 등급 타임라인 (평가 5회 + 현재)", det.locator(".tl-item").count() == 6, str(det.locator(".tl-item").count()))
    check("상세: 하락 표시와 변동 사유", "하락" in det.inner_text() and "납기 지연 증가" in det.inner_text())
    check("상세: 마지막 점은 '현재 · 시스템 계산'", "현재 · 시스템 계산" in det.locator(".tl-item").last.inner_text())
    check("상세: 공급 부품 분류 경로", "›" in det.locator(".supply-parts").inner_text())
    check("상세: 관련 직종 임금 상승률", "CNC 공작기계 조작원" in det.inner_text())
    page.screenshot(path=f"{SHOTS}/05_1920_detail.png")
    page.locator(".detail-toggle").click()
    check("아이콘으로 접기", page.locator("#selection-detail").count() == 0 and page.locator(".detail-toggle").get_attribute("aria-expanded") == "false")
    page.locator(".detail-toggle").click()
    check("아이콘으로 다시 펼치기", page.locator("#selection-detail").count() == 1)
    ask(page, "업체 F의 평가 근거를 정리해 줘")
    check("챗봇 결과가 나오면 선택 정보는 접히고 아이콘만 남음", page.locator("#selection-detail").count() == 0 and page.locator(".detail-toggle").count() == 1 and page.locator(".result-titles h3").count() == 1)
    page.locator(".detail-toggle").click()
    check("결과 위로 선택 정보 다시 열기", page.locator("#selection-detail").count() == 1 and page.locator(".result-titles h3").count() == 1)
    close_chat(page)

    # 25. 대금 결제 일정 + 알림
    bell = page.locator(".bell-badge").inner_text()
    check("헤더 알림 6건 (임박 5 + 지급 확인 필요 1)", bell == "6", bell)
    page.click(".bell")
    check("알림 → 구매조달 › 대금 결제 일정", page.locator(".subtab[data-sub='payments']").get_attribute("aria-selected") == "true")
    pv = page.locator(".subview")
    check("결제 일정: SAMPLE·기준일·출처 표시", "SAMPLE" in pv.locator(".sv-head").inner_text() and "기준일 2026-10-15" in pv.locator(".sv-head").inner_text())
    check("결제 일정: 임박 5건·예정 10건", pv.locator(".pay-card[data-pay='임박'] .pay-card-count").inner_text().startswith("5") and pv.locator(".pay-card[data-pay='예정'] .pay-card-count").inner_text().startswith("10"))
    check("결제 일정: 주별 그래프", pv.locator(".chart-card").count() == 1)
    row59 = pv.locator("#pay-table tbody tr", has_text="PO-2026-0059")
    check("PO-2026-0059: 월말 마감 익월 20일 → 2026-10-20 D-5 임박", "2026-10-20" in row59.inner_text() and "D-5" in row59.inner_text() and "임박" in row59.inner_text())
    pv.locator(".pay-card[data-pay='조건 확인 필요']").click()
    check("결제조건을 해석 못 하면 날짜를 추정하지 않음", pv.locator("#pay-table tbody tr").count() == 3 and "선적일" in pv.locator("#pay-table").inner_text())
    pv.locator(".pay-card[data-pay='임박']").click()
    check("상태 카드로 목록 거르기 (임박 5건)", pv.locator("#pay-table tbody tr").count() == 5)
    pv.locator("#pay-table tbody tr", has_text="PO-2026-0059").locator("input[type='checkbox']").check()
    check("'확인함' 표시하면 알림 건수 감소", page.locator(".bell-badge").inner_text() == "5")
    pv.locator(".pay-card[data-pay='임박']").click()
    n_s, bad_s = sort_all(page, "#pay-table")
    check(f"결제 일정 표: 정렬할 수 있는 모든 열({n_s}개) 두 방향 → 해제", n_s >= 8 and not bad_s, ", ".join(bad_s))
    page.screenshot(path=f"{SHOTS}/06_1920_payments.png")

    # 26. 원자재 시세
    page.locator(".tab", has_text="원가-계약").click()
    page.locator(".subtab[data-sub='materials']").click()
    mv = page.locator(".subview")
    check("원자재 시세: SAMPLE 표시·원/kg", "SAMPLE" in mv.locator(".sv-head").inner_text() and "원/kg" in mv.inner_text())
    check("원자재 시세: 7종 표", mv.locator(".mat-table tbody tr").count() == 7)
    mv.locator(".mat-chips .chip-btn", has_text="니켈").click()
    check("원자재 선택 → 국제·국내 추이 그래프", "니켈" in mv.locator(".chart-title").inner_text() and mv.locator(".legend-item").count() == 2)
    check("품목별 소재 ↔ 계약 단가 변동률", mv.locator(".link-table tbody tr").count() == 10 and "업체" in mv.locator(".link-table").inner_text())
    page.screenshot(path=f"{SHOTS}/07_1920_materials.png")

    # 27. 인건비 동향
    page.locator(".subtab[data-sub='wages']").click()
    wv = page.locator(".subview")
    check("인건비: 대분류 3개 트리", wv.locator(".wt-node.lv-0").count() == 3)
    wv.locator(".wt-fold[aria-label='금속 성형 기능직 펼치기']").click()
    check("중분류 펼치면 소분류 표시", wv.locator(".wt-name", has_text="주조원").count() == 1)
    wv.locator(".wt-name", has_text="주조원").click()
    check("소분류 선택 → 임금 추이·전년 대비", "주조원" in wv.locator(".chart-title").inner_text() and "전년 같은 달 대비" in wv.inner_text())
    check("인건비: 관련 협력사 표시", "업체 A" in wv.inner_text())

    # 28. 경쟁사 공급 동향 · 업계 동향
    page.locator(".tab", has_text="자재-공급관리").click()
    page.locator(".subtab[data-sub='competitors']").click()
    cv = page.locator(".subview")
    check("경쟁사: 가상 회사명·공급 경쟁 표시", cv.locator(".rel").count() == 7 and "(가상)" in cv.inner_text() and cv.locator(".rel.is-contested").count() == 4)
    check("경쟁사: 출처 링크 없음 경고", cv.locator(".src-missing").count() == 7)
    check("경쟁사: 기준일(가장 최근 확인일) 표시", "기준일 2026-" in cv.locator(".sv-head").inner_text() and "가장 최근 확인일" in cv.locator(".sv-head").inner_text())
    page.select_option("#comp-item", "ITM-1003")
    check("경쟁사: 부품으로 거르기", cv.locator(".rel").count() == 1)
    page.locator(".subtab[data-sub='news']").click()
    nv = page.locator(".subview")
    check("업계 동향: 기사 8건 (SAMPLE)", nv.locator(".news").count() == 8 and "SAMPLE" in nv.inner_text())
    page.select_option("#news-mat", "니켈")
    check("업계 동향: 원자재로 거르기 (니켈 2건)", nv.locator(".news").count() == 2)
    sf2 = small_fonts(page)
    check("새 화면도 글자 12px 이상", not sf2, ", ".join(sf2))
    page.locator(".subtab[data-sub='list']").click()

    # 29. 샘플 버튼 정리
    check("탭의 '샘플' 버튼 없음 (헤더 한 곳)", page.locator(".filter-row button", has_text="샘플").count() == 0 and page.locator(".topbar .sample-reset").count() == 1)
    page.click(".topbar .sample-reset")
    check("샘플로 되돌리기 → 확인 창", page.locator("[role='alertdialog']").count() == 1)
    inside = []
    for _ in range(4):
        page.keyboard.press("Tab")
        inside.append(page.evaluate("() => !!document.activeElement.closest('[role=alertdialog]')"))
    check("확인 창 안에서 Tab이 밖으로 나가지 않음", all(inside), str(inside))
    page.locator("[role='alertdialog'] .btn-quiet", has_text="취소").click()
    check("확인 창 닫으면 누른 버튼으로 초점", page.evaluate("() => new Promise(r => setTimeout(() => r(document.activeElement && document.activeElement.classList.contains('sample-reset')), 50))"))

    # 30. 업로드 → 저장 → 새로고침 유지 → 다시 분석 안내
    ask(page, "노즐 후보 업체 비교해 줘")
    close_chat(page)
    csv = "업체ID,협력사명,품질 평가 점수,불량률(%),기술·대응력 점수,경영·ESG 점수\nSUP-011,업체 K,60,4.5,50,40\n"
    path = os.path.join(SHOTS, "upload.csv")
    open(path, "w", encoding="utf-8").write(csv)
    page.locator(".topbar .btn-primary").click()
    page.select_option("#import-dataset", "suppliers")
    page.set_input_files("#import-file", path)
    page.wait_for_selector(".validation")
    page.locator(".modal-foot .btn-primary").click()
    page.wait_for_selector(".stale")
    check("데이터 변경 → 다시 분석 안내", page.locator(".stale").count() == 1)
    page.wait_for_selector(".ds-saved")
    page.wait_for_function("() => document.querySelector('.ds-saved').textContent.includes('저장됨')")
    check("헤더: 업로드 데이터 · 브라우저에 저장됨", "브라우저에 저장됨" in page.locator(".ds-saved").inner_text() and page.locator(".badge-demo").inner_text() == "업로드 포함")
    tb_h = page.locator(".topbar").bounding_box()["height"]
    check("업로드 상태에서도 1920 헤더 한 줄·가로 스크롤 없음", tb_h < 72 and no_hscroll(page) and page.locator(".ai-status-text").evaluate("e => e.getClientRects().length === 1"), str(round(tb_h)))
    page.locator(".stale button").click()
    page.wait_for_timeout(200)
    check("다시 분석 후 안내 해제", page.locator(".stale").count() == 0)
    page.reload(wait_until="domcontentloaded")
    page.wait_for_selector(".eval-card")
    check("새로고침 후 업로드 데이터 유지", page.locator(".badge-demo").inner_text() == "업로드 포함" and page.locator(".data-table tbody tr", has_text="업체 K").count() == 1)
    page.wait_for_selector(".result-titles h3")
    check("새로고침 후 '다시 분석' 안내가 잘못 뜨지 않음", page.locator(".stale").count() == 0)
    page.click(".user-btn")
    del_item = page.locator(".user-pop .menu-item", has_text="저장된 업로드 데이터 삭제")
    check("사용자 메뉴: 저장된 업로드 데이터 삭제 (업로드 상태에서 사용 가능)", del_item.is_enabled())
    del_item.click()
    check("저장된 업로드 데이터 삭제 → 확인 창", page.locator("[role='alertdialog']").count() == 1 and "업로드" in page.locator("[role='alertdialog']").inner_text())
    page.locator("[role='alertdialog'] [data-confirm]").click()
    check("샘플로 되돌리기 → SAMPLE", page.locator(".badge-demo").inner_text() == "SAMPLE")
    page.reload(wait_until="domcontentloaded")
    page.wait_for_selector(".eval-card")
    check("새로고침해도 샘플 유지 (저장본 삭제됨)", page.locator(".badge-demo").inner_text() == "SAMPLE")
    page.click(".user-btn")
    check("SAMPLE 상태에서는 '저장된 업로드 데이터 삭제' 비활성", page.locator(".user-pop .menu-item", has_text="저장된 업로드 데이터 삭제").is_disabled())
    page.keyboard.press("Escape")
    check("콘솔 오류 없음", not errors, "; ".join(errors[:3]))

    # 31. 로그아웃 · 사용자별 저장 분리
    n_hist_hong = int(page.locator(".history-toggle .count-badge").inner_text())
    page.click(".user-btn")
    page.locator(".user-pop .menu-item", has_text="로그아웃").click()
    page.locator("[role='alertdialog'] [data-confirm]").click()
    page.wait_for_selector("#login-name")
    check("로그아웃 → 로그인 화면 (최근 사용자 표시)", "홍길동" in page.locator(".login-recent").inner_text())
    page.fill("#login-name", "김철수")
    page.click("#login-submit")
    page.click("#ob-skip")
    page.wait_for_selector(".eval-card")
    check("다른 사용자: 분석 이력·관심 목록 따로", page.locator(".history-toggle .count-badge").inner_text() == "0" and page.locator(".fav-star.is-on").count() == 0)
    page.click(".user-btn")
    page.locator(".user-pop .menu-item", has_text="로그아웃").click()
    page.locator("[role='alertdialog'] [data-confirm]").click()
    page.locator(".login-recent .chip-btn", has_text="홍길동").click()
    page.wait_for_selector(".eval-card")
    check("원래 사용자로 돌아오면 이력·관심 목록 그대로", int(page.locator(".history-toggle .count-badge").inner_text()) == n_hist_hong and page.locator(".fav-star.is-on").count() >= 1)
    hong_id = page.evaluate("() => (JSON.parse(localStorage.getItem('procure-ai:users.v1')) || []).find(u => u.name === '홍길동').id")
    check("홍길동 기록이 이 브라우저에 있음", page.evaluate("(id) => Object.keys(localStorage).some(k => k.includes(':u:' + id + ':'))", hong_id))
    page.click(".user-btn")
    page.locator(".user-pop .menu-item", has_text="로그아웃").click()
    page.check("#logout-wipe")
    page.locator("[role='alertdialog'] [data-confirm]").click()
    page.wait_for_selector("#login-name")
    page.wait_for_timeout(800)  # 늦게 도는 저장이 지운 기록을 되살리지 않는지 확인
    left = page.evaluate("(id) => Object.keys(localStorage).filter(k => k.includes(':u:' + id + ':'))", hong_id)
    check("로그아웃 시 내 데이터 지우기 (최근 사용자 목록·저장 기록 삭제, 다른 사용자 기록은 남음)", "홍길동" not in (page.locator(".login-recent").inner_text() if page.locator(".login-recent").count() else "") and not left and "김철수" in page.locator(".login-recent").inner_text(), ", ".join(left))
    page.screenshot(path=f"{SHOTS}/08_1920_login.png")

    # 32. 1366×768
    lap = browser.new_page(viewport={"width": 1366, "height": 768})
    demo_login(lap)
    check("1366: 가로 스크롤 없음", no_hscroll(lap))
    check("1366: 상단 바 한 줄", lap.locator(".topbar").bounding_box()["height"] < 72, str(round(lap.locator(".topbar").bounding_box()["height"])))
    check("1366: 채점 카드 높이 줄어듦 (처리 단계 영역 없음)", lap.locator(".eval-card").bounding_box()["height"] < 360, str(round(lap.locator(".eval-card").bounding_box()["height"])))
    lb2 = lap.locator(".chat-launcher").bounding_box()
    check("1366: 챗봇 아이콘 왼쪽 아래", lb2["x"] < 60 and lb2["y"] > 768 - 100)
    sf3 = small_fonts(lap)
    check("1366: 글자 12px 이상", not sf3, ", ".join(sf3))
    lap.screenshot(path=f"{SHOTS}/09_1366.png")

    # 33. 태블릿 900
    tab = browser.new_page(viewport={"width": 900, "height": 1180})
    demo_login(tab)
    check("태블릿: 데이터 관리/분석 결과 탭 2개", tab.locator(".panel-tabs button").count() == 2 and tab.locator(".panel-a").is_visible() and not tab.locator(".panel-c").is_visible())
    tab.click(".chat-launcher")
    tab.wait_for_selector(".chat-window.is-sheet")
    check("태블릿: 챗봇은 하단 시트", tab.locator(".chat-window.is-sheet").count() == 1)
    tab.fill("#chat-input", "노즐 후보 업체 비교해 줘")
    tab.press("#chat-input", "Enter")
    tab.wait_for_selector(".result-table")
    check("태블릿: 분석 후 결과 탭으로 이동", tab.locator(".panel-c").is_visible())
    tab.locator(".chat-tools button[aria-label='닫기']").click()
    tab.screenshot(path=f"{SHOTS}/10_tablet.png")
    check("태블릿: 가로 스크롤 없음", no_hscroll(tab))
    sf5 = small_fonts(tab)
    check("태블릿: 글자 12px 이상", not sf5, ", ".join(sf5))
    tab.locator(".panel-tabs button", has_text="데이터 관리").click()
    tab.locator(".data-table tbody tr", has_text="업체 B").locator(".name").click()
    check("태블릿: 행을 누르면 분석 결과 탭으로 바뀌고 선택 정보 표시", tab.locator(".panel-c").is_visible() and tab.locator("#selection-detail").count() == 1)

    # 34. 휴대폰 400
    ph = browser.new_page(viewport={"width": 400, "height": 820})
    demo_login(ph)
    check("휴대폰: 가로 스크롤 없음", no_hscroll(ph))
    check("휴대폰: SAMPLE 표시 유지", ph.locator(".badge-demo").is_visible())
    check("휴대폰: 챗봇 버튼에 'AI 질문' 글자 (아이콘만 두지 않음)", "AI 질문" in ph.locator(".chat-launcher").inner_text())
    check("휴대폰: 상단 바 두 줄 (브랜드·알림·사용자 / 데이터 상태·불러오기)", ph.locator(".user-btn").bounding_box()["y"] < ph.locator(".topbar .btn-primary").bounding_box()["y"])
    sf4 = small_fonts(ph)
    check("휴대폰: 글자 12px 이상", not sf4, ", ".join(sf4))
    ph.screenshot(path=f"{SHOTS}/11_phone.png")
    browser.close()

failed = [r for r in results if not r[1]]
print(f"\n{len(results) - len(failed)}/{len(results)} 통과 · 스크린샷: {SHOTS}")
sys.exit(1 if failed else 0)
